"""Bounded, read-only tabular imports. This module never runs uploaded code.

Values stay strings; choosing numeric columns and decimal conventions belongs to
the dataset configuration step. ``header_row`` counts parsed rows from one.
"""

from __future__ import annotations

from .localization import tr


import csv
from collections import Counter
from datetime import date, datetime, time
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import PurePath
import re
import struct
from typing import Any
import zipfile
from xml.etree import ElementTree


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 75 * 1024 * 1024
MAX_ZIP_MEMBERS = 2000
MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_TABLES = 30
MAX_CELL_CHARS = 2000
PREVIEW_ROWS = 6


class ImportError(ValueError):
    """A readable import error safe to present in the interface."""


def _cell(value: Any) -> str:
    if value is None:
        result = ""
    elif isinstance(value, (datetime, date, time)):
        result = value.isoformat()
    elif isinstance(value, bool):
        result = "true" if value else "false"
    elif isinstance(value, float):
        result = str(int(value)) if value.is_integer() else str(value)
    else:
        result = str(value)
    if len(result) > MAX_CELL_CHARS:
        raise ImportError(f"{tr('A cell contains more than ')}{MAX_CELL_CHARS}{tr(' characters.')}")
    if "\x00" in result:
        raise ImportError(tr('A null byte was found. Check the file format and encoding.'))
    return result


def _check_row(row: list[str], count: int) -> None:
    # One additional row is allowed for the column headings.
    if count > MAX_ROWS + 1:
        raise ImportError(f"{tr('The table exceeds ')}{MAX_ROWS}{tr(' data rows.')}")
    if len(row) > MAX_COLUMNS:
        raise ImportError(f"{tr('The table exceeds ')}{MAX_COLUMNS}{tr(' columns.')}")


def _decode(content: bytes, encoding: str | None = None) -> tuple[str, list[str]]:
    aliases = {"utf-8": "utf-8-sig", "utf8": "utf-8-sig", "utf-8-sig": "utf-8-sig",
               "utf-16": "utf-16", "utf16": "utf-16", "cp1251": "cp1251", "windows-1251": "cp1251"}
    warnings = []
    if encoding is not None:
        if not isinstance(encoding, str) or encoding.lower() not in aliases:
            raise ImportError(tr('Supported encodings: UTF-8, UTF-16 with BOM and Windows-1251.'))
        selected = aliases[encoding.lower()]
    elif content.startswith((b"\xff\xfe", b"\xfe\xff")):
        selected = "utf-16"
    else:
        selected = "utf-8-sig"
    try:
        text = content.decode(selected, errors="strict")
    except (UnicodeError, LookupError) as exc:
        if encoding is not None or selected == "utf-16":
            raise ImportError(tr('Cannot read the file in the selected encoding.')) from exc
        try:
            text = content.decode("cp1251", errors="strict")
            selected = "cp1251"
        except UnicodeError as fallback_exc:
            raise ImportError(tr('Cannot detect encoding. Save the file as UTF-8.')) from fallback_exc
    if selected == "cp1251":
        warnings.append(tr('Windows-1251 encoding was used. Review column names and values.'))
    if "\x00" in text:
        raise ImportError(tr('Null bytes prevent reading. UTF-16 requires a BOM.'))
    return text, warnings


def _csv_rows(text: str, delimiter: str) -> list[list[str]]:
    rows = []
    try:
        for row in csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True):
            values = [_cell(value) for value in row]
            _check_row(values, len(rows) + 1)
            rows.append(values)
    except csv.Error as exc:
        raise ImportError(tr('Invalid CSV: check quotes and line separators.')) from exc
    return rows


def _delimiter(text: str, preferred: str) -> str:
    scores = []
    # csv.reader respects quoted separators and quoted multiline cells.
    for candidate in (preferred, *[d for d in (",", ";", "\t") if d != preferred]):
        widths = []
        try:
            reader = csv.reader(StringIO(text, newline=""), delimiter=candidate, strict=True)
            for index, row in enumerate(reader):
                if row and any(value.strip() for value in row):
                    widths.append(len(row))
                if index >= 99:
                    break
        except csv.Error:
            # Keep evidence from a valid header; the real parse must expose a
            # malformed later row instead of guessing a one-column format.
            pass
        counts = Counter(widths)
        multi = [(count, width) for width, count in counts.items() if width > 1]
        if multi:
            count, width = max(multi)
            scores.append(((count / max(1, len(widths)), count, width), candidate))
    if not scores:
        return preferred
    best_score = max(score for score, _ in scores)
    return next(candidate for score, candidate in scores if score == best_score)


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[dict] = []
        self.table: dict | None = None
        self.row: list[str] | None = None
        self.parts: list[str] | None = None
        self.cell_length = 0
        self.caption = False
        self.caption_parts: list[str] = []
        self.ignored: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "template"):
            self.ignored.append(tag)
            return
        if self.ignored:
            return
        if tag == "table":
            if self.table is not None:
                raise ImportError(tr('Nested HTML tables are unsupported. Export a flat table.'))
            if len(self.tables) >= MAX_TABLES:
                raise ImportError(f"{tr('The file exceeds ')}{MAX_TABLES}{tr(' tables.')}")
            self.table = {"id": f"html:{len(self.tables)}", "label": f"{tr('Table ')}{len(self.tables) + 1}",
                          "rows": [], "warnings": []}
        elif self.table is not None:
            if tag == "caption":
                self.caption = True
                self.caption_parts = []
            elif tag == "tr":
                if self.row is not None or self.parts is not None:
                    raise ImportError(tr('Invalid HTML table: unclosed rows or cells.'))
                self.row = []
            elif tag in ("td", "th"):
                if self.row is None or self.parts is not None:
                    raise ImportError(tr('Invalid HTML cell structure.'))
                attributes = dict(attrs)
                for span in ("rowspan", "colspan"):
                    try:
                        if int(attributes.get(span, "1")) != 1:
                            raise ValueError()
                    except (TypeError, ValueError):
                        raise ImportError(tr('HTML rowspan/colspan cells are unsupported. Export a flat table.')) from None
                self.parts = []
                self.cell_length = 0
            elif tag == "br" and self.parts is not None:
                self.parts.append(" ")

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag != "br":
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.ignored:
            if tag == self.ignored[-1]:
                self.ignored.pop()
            return
        if self.table is None:
            return
        if tag in ("td", "th"):
            if self.parts is None or self.row is None:
                raise ImportError(tr('Invalid HTML table closing tags.'))
            self.row.append(_cell("".join(self.parts).strip()))
            _check_row(self.row, len(self.table["rows"]) + 1)
            self.parts = None
        elif tag == "tr":
            if self.row is None or self.parts is not None:
                raise ImportError(tr('Invalid or unclosed HTML table row.'))
            self.table["rows"].append(self.row)
            _check_row(self.row, len(self.table["rows"]))
            self.row = None
        elif tag == "caption":
            self.caption = False
            caption = "".join(self.caption_parts).strip()
            if caption:
                self.table["label"] = _cell(caption)
        elif tag == "table":
            if self.row is not None or self.parts is not None:
                raise ImportError(tr('HTML-table contains an unclosed row or cell.'))
            self.tables.append(self.table)
            self.table = None

    def handle_data(self, data):
        if self.ignored:
            return
        if self.parts is not None:
            self.cell_length += len(data)
            if self.cell_length > MAX_CELL_CHARS:
                raise ImportError(f"{tr('A cell contains more than ')}{MAX_CELL_CHARS}{tr(' characters.')}")
            self.parts.append(data)
        elif self.caption:
            self.caption_parts.append(data)
            if sum(map(len, self.caption_parts)) > MAX_CELL_CHARS:
                raise ImportError(tr('HTML table name is too long.'))


def _zip_validate(content: bytes) -> dict[str, bool]:
    merged = {}
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise ImportError(tr('Too many parts in the Excel file.'))
            if sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
                raise ImportError(tr('Expanded Excel exceeds 75 MiB.'))
            if len({member.filename for member in members}) != len(members):
                raise ImportError(tr('Excel contains duplicate archive members.'))
            for member in members:
                if member.flag_bits & 1:
                    raise ImportError(tr('Encrypted Excel is unsupported. Save a copy without a password.'))
                if "vbaproject" in member.filename.lower():
                    raise ImportError(tr('Macro-enabled workbooks are unsupported. Save without macros.'))
                if not member.filename.lower().endswith((".xml", ".rels")):
                    continue
                data = archive.read(member)
                if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.IGNORECASE):
                    raise ImportError(tr('XML with DTD or entity declarations is unsupported.'))
                if member.filename.startswith("xl/worksheets/"):
                    has_merge = False
                    for _, element in ElementTree.iterparse(BytesIO(data), events=("end",)):
                        if element.tag.rsplit("}", 1)[-1] == "mergeCell":
                            has_merge = True
                        element.clear()
                    merged[member.filename] = has_merge
    except (zipfile.BadZipFile, ElementTree.ParseError, OSError, RuntimeError) as exc:
        raise ImportError(tr('Corrupt Excel XLSX file.')) from exc
    return merged


def _excel_value(value: Any, number_format: str = "") -> str:
    # A numeric cell formatted as 00000 is frequently an identifier.
    if isinstance(value, (int, float)) and not isinstance(value, bool) and re.fullmatch(r"0{2,}", number_format):
        if float(value).is_integer() and value >= 0:
            return _cell(str(int(value)).zfill(len(number_format)))
    return _cell(value)


def _xlsx(content: bytes) -> list[dict]:
    merged = _zip_validate(content)
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise ImportError(tr('Excel library is missing. Restart using start_residual.')) from exc
    book = None
    try:
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) > MAX_TABLES:
            raise ImportError(f"{tr('The workbook exceeds ')}{MAX_TABLES}{tr(' sheets.')}")
        tables = []
        for index, sheet in enumerate(book.worksheets):
            if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 1:
                raise ImportError(f"{tr('Sheet «')}{sheet.title}{tr('» exceeds the limit of ')}{MAX_ROWS}{tr(' rows or ')}{MAX_COLUMNS}{tr(' columns.')}")
            # Some exporters write A1:A1 even when the sheet contains more
            # data. Never trust that dimension and silently drop later cells.
            sheet.reset_dimensions()
            rows = []
            has_formula = False
            has_error = False
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    has_formula |= cell.data_type == "f"
                    has_error |= cell.data_type == "e"
                    values.append(_excel_value(cell.value, cell.number_format))
                _check_row(values, len(rows) + 1)
                rows.append(values)
            # Drop merely formatted empty rightmost columns and trailing rows.
            rows = _trim_excel(rows)
            warnings = []
            blockers = []
            if has_formula:
                blockers.append(tr('The sheet contains formulas. Save a values-only copy; import does not recalculate formulas.'))
            if has_error:
                blockers.append(tr('The sheet contains Excel errors. Fix them before import.'))
            path = getattr(sheet, "_worksheet_path", "").lstrip("/")
            if path not in merged:
                raise ImportError(tr('Cannot validate the Excel sheet structure.'))
            if merged[path]:
                blockers.append(tr('Merged cells found. Unmerge them and save a flat table.'))
            tables.append({"id": f"sheet:{index}", "label": sheet.title, "rows": rows,
                           "warnings": warnings + blockers, "blockers": blockers})
        return tables
    except ImportError:
        raise
    except Exception as exc:
        raise ImportError(tr('Cannot read XLSX. Save a regular Excel workbook without a password.')) from exc
    finally:
        if book is not None:
            book.close()


def _trim_excel(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(value.strip() for value in rows[-1]):
        rows.pop()
    width = max((index + 1 for row in rows for index, value in enumerate(row) if value.strip()), default=0)
    return [(row[:width] + [""] * max(0, width - len(row))) for row in rows]


def _xls_formula(book, index: int) -> bool:
    # xlrd intentionally exposes cached values, so inspect BIFF cell records as
    # well. Restrict this adapter to BIFF5/8, whose sheet offsets are available.
    if book.biff_version < 50:
        raise ImportError(tr('This old XLS variant is unsupported. Save as XLSX or CSV.'))
    offset = book._sh_abs_posn[index]
    memory = book.mem
    if memory is None:
        raise ImportError(tr('Cannot validate formulas in XLS.'))
    while offset + 4 <= len(memory):
        record, size = struct.unpack_from("<HH", memory, offset)
        offset += 4
        if offset + size > len(memory):
            raise ImportError(tr('Corrupt XLS structure.'))
        if record in (0x0006, 0x0206, 0x0406):
            return True
        if record == 0x000A:
            return False
        offset += size
    raise ImportError(tr('Cannot validate the end of the XLS sheet.'))


def _xls(content: bytes) -> list[dict]:
    try:
        import xlrd
    except ModuleNotFoundError as exc:
        raise ImportError(tr('XLS library is missing. Restart using start_residual.')) from exc
    book = None
    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True, logfile=StringIO())
        if book.nsheets > MAX_TABLES:
            raise ImportError(f"{tr('The workbook exceeds ')}{MAX_TABLES}{tr(' sheets.')}")
        tables = []
        for index in range(book.nsheets):
            has_formula = _xls_formula(book, index)
            sheet = book.sheet_by_index(index)
            if sheet.nrows > MAX_ROWS + 1 or sheet.ncols > MAX_COLUMNS:
                raise ImportError(f"{tr('Sheet «')}{sheet.name}{tr('» exceeds the limit of ')}{MAX_ROWS}{tr(' rows or ')}{MAX_COLUMNS}{tr(' columns.')}")
            rows = []
            has_error = False
            for row_index in range(sheet.nrows):
                values = []
                for cell in sheet.row(row_index):
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        has_error = True
                        value = xlrd.error_text_from_code.get(value, tr('Excel error'))
                    format_key = book.xf_list[cell.xf_index].format_key
                    number_format = book.format_map[format_key].format_str
                    values.append(_excel_value(value, number_format))
                _check_row(values, len(rows) + 1)
                rows.append(values)
            blockers = []
            if has_formula:
                blockers.append(tr('XLS sheet contains formulas. Save values only; import does not recalculate.'))
            if has_error:
                blockers.append(tr('The sheet contains Excel errors. Fix them before import.'))
            if sheet.merged_cells:
                blockers.append(tr('Merged cells found. Unmerge them and save a flat table.'))
            tables.append({"id": f"sheet:{index}", "label": sheet.name, "rows": _trim_excel(rows),
                           "warnings": list(blockers), "blockers": blockers})
            book.unload_sheet(index)
        return tables
    except ImportError:
        raise
    except Exception as exc:
        raise ImportError(tr('Cannot read XLS. Save an unprotected workbook or export CSV.')) from exc
    finally:
        if book is not None:
            book.release_resources()


def _load(name: str, content: bytes, options: dict | None = None) -> tuple[str, list[dict], list[str]]:
    if not isinstance(name, str) or not name.strip():
        raise ImportError(tr('Provide a filename with an extension.'))
    if not isinstance(content, bytes) or not content:
        raise ImportError(tr('File is empty or missing.'))
    if len(content) > MAX_FILE_BYTES:
        raise ImportError(tr('File exceeds the 10 MiB limit.'))
    if options is None:
        options = {}
    if not isinstance(options, dict) or set(options) - {"encoding", "delimiter"}:
        raise ImportError(tr('Unknown import settings.'))
    extension = PurePath(name).suffix.lower()
    supported = {".csv", ".tsv", ".xlsx", ".xls", ".html", ".htm"}
    if extension not in supported:
        raise ImportError(tr('CSV, TSV, XLSX, XLS and HTML tables are supported. Python and other programs are never executed.'))
    warnings = []
    if extension in (".csv", ".tsv", ".html", ".htm"):
        text, warnings = _decode(content, options.get("encoding"))
        if extension in (".csv", ".tsv"):
            separator = options.get("delimiter")
            if separator is not None and separator not in (",", ";", "\t"):
                raise ImportError(tr('Delimiter must be comma, semicolon or tab.'))
            separator = separator or _delimiter(text, "\t" if extension == ".tsv" else ",")
            rows = _csv_rows(text, separator)
            tables = [{"id": "csv:0", "label": PurePath(name).name, "rows": rows, "warnings": []}]
            format_name = extension[1:]
        else:
            parser = _Tables()
            try:
                parser.feed(text)
                parser.close()
            except ImportError:
                raise
            except Exception as exc:
                raise ImportError(tr('Cannot read the HTML table.')) from exc
            if parser.table is not None:
                raise ImportError(tr('HTML-table is not closed. Export valid HTML or CSV.'))
            tables, format_name = parser.tables, "html"
            warnings.append(tr('Only HTML tables are read. Scripts, styles and external resources are not run or downloaded.'))
    elif extension == ".xlsx":
        tables, format_name = _xlsx(content), "xlsx"
    else:
        tables, format_name = _xls(content), "xls"
    if not tables or not any(table["rows"] for table in tables):
        raise ImportError(tr('No nonempty tables found.'))
    return format_name, tables, warnings


def inspect_file(name: str, content: bytes) -> dict:
    """List tables and raw previews. Header and data validity are checked on read."""
    format_name, tables, warnings = _load(name, content)
    return {"format": format_name, "tables": [
        {"id": table["id"], "label": table["label"],
         "row_count": sum(bool(any(value.strip() for value in row)) for row in table["rows"][1:]),
         "preview": table["rows"][:PREVIEW_ROWS], "warnings": table["warnings"]}
        for table in tables], "warnings": warnings}


def read_table(name: str, content: bytes, table_id: str, header_row: int = 1, options: dict | None = None) -> dict:
    """Read one rectangular table without type inference or formula execution."""
    format_name, tables, warnings = _load(name, content, options)
    table = next((table for table in tables if table["id"] == table_id), None)
    if table is None:
        raise ImportError(tr('Selected table or sheet not found.'))
    if table.get("blockers"):
        raise ImportError(" ".join(table["blockers"]))
    raw = table["rows"]
    if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= len(raw):
        raise ImportError(tr('Choose an existing header row, starting at 1.'))
    columns = [value.strip() for value in raw[header_row - 1]]
    if not columns or any(not value for value in columns):
        raise ImportError(tr('Header contains empty column names. Choose another row or fix the file.'))
    if len({value.casefold() for value in columns}) != len(columns):
        raise ImportError(tr('Column names repeat. Give each column a unique name.'))
    rows = []
    blank_count = 0
    for index, row in enumerate(raw[header_row:], start=header_row + 1):
        if not row or not any(value.strip() for value in row):
            blank_count += 1
            continue
        if len(row) != len(columns):
            raise ImportError(f"{tr('Row ')}{index}{tr(' contains ')}{len(row)}{tr(' cells instead of ')}{len(columns)}{tr('. Check delimiters and table structure.')}")
        rows.append(row)
    if not rows:
        raise ImportError(tr('No data after the selected header row.'))
    if blank_count:
        warnings.append(f"{tr('Completely empty rows skipped: ')}{blank_count}.")
    return {"columns": columns, "rows": rows, "warnings": warnings + table["warnings"],
            "table_label": table["label"], "format": format_name, "header_row": header_row}
